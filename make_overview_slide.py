# Build a FINRA-branded PowerPoint deck summarizing the enterprise & emerging
# risk pipeline. Two slides are appended onto the official FINRA template
# (tests/PPT_This_Is_FINRA_&_FF_Wht_Template.potx) so they inherit the real
# master, fonts, logo, and footer:
#   1. Technical overview  — the 8-stage pipeline + input lanes (live numbers).
#   2. Plain-language slide — the same process for a non-technical reader.
#
#   python make_overview_slide.py
#
# The template is a .potx (PowerPoint template). python-pptx won't open that
# content-type directly, so we transparently re-stamp it as a .pptx first
# (see load_template) — the XML is otherwise identical.

import os
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# --- FINRA brand palette (matches the template theme exactly) ----------------
CORE   = RGBColor(0x23, 0x3E, 0x66)
ACCENT = RGBColor(0x00, 0x82, 0xD1)
GRAY   = RGBColor(0x59, 0x59, 0x59)
GREEN  = RGBColor(0x9E, 0xC4, 0x05)
YELLOW = RGBColor(0xFF, 0xCF, 0x40)
RED    = RGBColor(0xFB, 0x48, 0x3D)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
BODY   = RGBColor(0x33, 0x33, 0x33)
MUTED  = RGBColor(0x6B, 0x72, 0x80)
PANEL  = RGBColor(0xFF, 0xFF, 0xFF)
LINE   = RGBColor(0xE2, 0xE6, 0xEC)
BG     = RGBColor(0xF4, 0xF6, 0xF9)

FONT = "Open Sans"

TEMPLATE = Path("tests/PPT_This_Is_FINRA_&_FF_Wht_Template.potx")
OUT = Path("output/risk_pipeline_overview.pptx")

# "Headline Only" layout: branded master chrome (logo/footer) + a top Title
# placeholder, with the whole body area free for our custom diagram.
HEADLINE_ONLY_LAYOUT = 3

# Title sits at ~0.53in tall; start body content below it.
BODY_TOP = Inches(1.45)


def load_template():
    """Open the FINRA .potx as an editable Presentation.

    python-pptx rejects the template content-type, so copy the archive to a
    temp .pptx, rewriting only the [Content_Types].xml declaration. If the
    template is missing, fall back to a blank 16:9 deck so the script still runs.
    """
    if not TEMPLATE.exists():
        print(f"NOTE: template {TEMPLATE} not found; using a blank deck.")
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        return prs

    OUT.parent.mkdir(exist_ok=True)
    tmp = OUT.parent / "_template_work.pptx"
    with zipfile.ZipFile(TEMPLATE) as zin, \
            zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for n in zin.namelist():
            data = zin.read(n)
            if n == "[Content_Types].xml":
                data = data.decode("utf-8").replace(
                    "presentationml.template.main+xml",
                    "presentationml.presentation.main+xml").encode("utf-8")
            zout.writestr(n, data)
    prs = Presentation(str(tmp))
    os.remove(tmp)
    return prs


def _stats(risk_type):
    """Live (risks, articles, topics) for a risk type, from the output CSVs."""
    import pandas as pd
    risks = articles = topics = None
    sent = Path(f"output/{risk_type}_risks_online_sentiment.csv")
    summ = Path(f"output/{risk_type}_risks_topic_summary.csv")
    if sent.exists():
        d = pd.read_csv(sent)
        articles = len(d)
        risks = int(d["RISK_ID"].nunique()) if "RISK_ID" in d else None
    if summ.exists():
        topics = len(pd.read_csv(summ))
    return risks, articles, topics


def _fmt_articles(n):
    if n is None:
        return "—"
    if n >= 1000:
        return f"~{round(n / 100) * 100:,}"
    return f"{n:,}"


def _slide_dims(prs):
    return prs.slide_width, prs.slide_height


def _add_content_slide(prs, title):
    """Add a slide on the branded 'Headline Only' layout and set its title."""
    try:
        layout = prs.slide_layouts[HEADLINE_ONLY_LAYOUT]
    except IndexError:
        layout = prs.slide_layouts[-1]
    slide = prs.slides.add_slide(layout)
    # fill the title placeholder if present
    set_title = False
    for ph in slide.placeholders:
        if ph.placeholder_format.type == 1 or ph.placeholder_format.idx == 0:  # TITLE
            ph.text = title
            set_title = True
            break
    if not set_title:
        _text(slide, Inches(0.67), Inches(0.53), Inches(10.75), Inches(0.7),
              [[(title, 24, CORE, True, False)]])
    return slide


def _box(slide, x, y, w, h, fill=None, line=None, line_w=0.75, rounded=False):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, x, y, w, h)
    if rounded:
        try:
            shp.adjustments[0] = 0.08
        except Exception:
            pass
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(line_w)
    shp.shadow.inherit = False
    return shp


def _text(slide, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
          space_after=2, line_spacing=1.0):
    """runs: list of paragraphs; each is a list of (text,size,color,bold,italic)."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Pt(2)
    tf.margin_top = tf.margin_bottom = Pt(1)
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        p.space_before = Pt(0)
        p.line_spacing = line_spacing
        for (txt, size, color, bold, italic) in para:
            r = p.add_run()
            r.text = txt
            r.font.size = Pt(size)
            r.font.color.rgb = color
            r.font.bold = bold
            r.font.italic = italic
            r.font.name = FONT
    return tb


def build_technical_slide(prs):
    slide = _add_content_slide(prs, "Enterprise & Emerging Risk Signal Pipeline")
    slide_w, _ = _slide_dims(prs)
    margin = Inches(0.67)
    total_w = slide_w - 2 * margin

    _text(slide, margin, Inches(1.28), total_w, Inches(0.35),
          [[("How news coverage becomes tracked risk signals", 13, ACCENT, True, False)]])

    # --- pipeline flow: 8 stages as connected cards -------------------------
    stages = [
        ("1  Fetch", "Pull news per risk\nsearch term (newsdata.io)"),
        ("2  Sentiment", "VADER scores each\narticle +/-/neutral"),
        ("3  spaCy NLP", "Entities & keyword\nnoun-phrases"),
        ("4  Cluster", "Embed + dedup +\nBERTopic topics"),
        ("5  Quality", "Score 0-1, drop\nboilerplate / wire"),
        ("6  Trends", "Monthly signal\nintensity per topic"),
        ("7  Gaps", "Taxonomy-fit &\nnovelty scoring"),
        ("8  Report", "Interactive HTML\nexplorers + index"),
    ]
    n = len(stages)
    gap = Inches(0.10)
    flow_top = Inches(1.95)
    flow_h = Inches(1.35)
    card_w = Emu(int((total_w - gap * (n - 1)) / n))
    accent_cycle = [ACCENT, GREEN, ACCENT, CORE, GREEN, ACCENT, RED, CORE]
    for i, (title, desc) in enumerate(stages):
        x = Emu(int(margin + i * (card_w + gap)))
        _box(slide, x, flow_top, card_w, flow_h, fill=PANEL, line=LINE, line_w=1.0, rounded=True)
        _box(slide, x, flow_top, card_w, Pt(5), fill=accent_cycle[i % len(accent_cycle)])
        _text(slide, x + Inches(0.08), flow_top + Inches(0.12), card_w - Inches(0.16), Inches(0.4),
              [[(title, 12.5, CORE, True, False)]])
        _text(slide, x + Inches(0.08), flow_top + Inches(0.52), card_w - Inches(0.16), Inches(0.8),
              [[(ln, 9, GRAY, False, False)] for ln in desc.split("\n")], line_spacing=1.05)
        if i < n - 1:
            arrow = slide.shapes.add_shape(
                MSO_SHAPE.RIGHT_ARROW,
                Emu(int(x + card_w - Inches(0.02))), flow_top + Inches(0.55),
                Inches(0.14), Inches(0.22))
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = MUTED
            arrow.line.fill.background()
            arrow.shadow.inherit = False

    # --- two input lanes ----------------------------------------------------
    ent_risks, ent_art, ent_top = _stats("enterprise")
    emg_risks, emg_art, emg_top = _stats("emerging")
    lane_top = Inches(3.65)
    lane_h = Inches(1.6)
    half = Emu(int((total_w - Inches(0.3)) / 2))

    _box(slide, margin, lane_top, half, lane_h, fill=PANEL, line=LINE, line_w=1.0, rounded=True)
    _box(slide, margin, lane_top, Inches(0.09), lane_h, fill=CORE)
    _text(slide, margin + Inches(0.25), lane_top + Inches(0.14), half - Inches(0.4), Inches(1.4),
          [[("Enterprise risks", 15, CORE, True, False)],
           [(f"{ent_risks} risks", 11, ACCENT, True, False), ("  ·  keyword search terms from ", 10.5, GRAY, False, False),
            ("EnterpriseRisksListEncoded.csv", 9.5, MUTED, False, True)],
           [(f"{_fmt_articles(ent_art)} articles", 11, CORE, True, False), (" clustered into ", 10.5, GRAY, False, False),
            (f"{ent_top} topics", 11, CORE, True, False)],
           [("Known risks already on the taxonomy — tracked over time.", 10, GRAY, False, True)]],
          line_spacing=1.08, space_after=4)

    ex = Emu(int(margin + half + Inches(0.3)))
    _box(slide, ex, lane_top, half, lane_h, fill=PANEL, line=LINE, line_w=1.0, rounded=True)
    _box(slide, ex, lane_top, Inches(0.09), lane_h, fill=GREEN)
    _text(slide, ex + Inches(0.25), lane_top + Inches(0.14), half - Inches(0.4), Inches(1.4),
          [[("Emerging risks", 15, CORE, True, False)],
           [(f"{emg_risks} risks", 11, ACCENT, True, False), ("  ·  keyword search terms from ", 10.5, GRAY, False, False),
            ("EmergingRisksListEncoded.csv", 9.5, MUTED, False, True)],
           [(f"{_fmt_articles(emg_art)} articles", 11, CORE, True, False), (" clustered into ", 10.5, GRAY, False, False),
            (f"{emg_top} topics", 11, CORE, True, False)],
           [("Forward-looking risks watched for early signal growth.", 10, GRAY, False, True)]],
          line_spacing=1.08, space_after=4)

    _text(slide, margin, Inches(5.4), total_w, Inches(0.55),
          [[("Both risk types run the identical 8-stage pipeline; only the search-term list differs. ",
             9.5, MUTED, False, False),
            ("Stages 1-2: news_sentiment_scraper.py · 3-5: topic_clustering.py · 6: topic_trends.py · 7: topic_gaps.py · 8: build_report.py + build_index.py",
             8.5, MUTED, False, True)]],
          line_spacing=1.05)


def build_plain_language_slide(prs):
    slide = _add_content_slide(prs, "In Plain Language: How It Works")
    slide_w, _ = _slide_dims(prs)
    margin = Inches(0.67)
    total_w = slide_w - 2 * margin

    _text(slide, margin, Inches(1.28), total_w, Inches(0.55),
          [[("Think of it as a tireless news-reader for risk: ", 13.5, CORE, True, False),
            ("it reads the news every day, sorts it into themes, and flags what's getting "
             "louder or looks new — so people can act sooner.", 13.5, BODY, False, False)]],
          line_spacing=1.1)

    steps = [
        ("Gather the news",
         "Every day it collects news articles about the risks we care about, from reputable outlets."),
        ("Read the mood",
         "It judges whether each story is positive, negative, or neutral — tone, not just volume."),
        ("Group similar stories",
         "It bundles articles about the same thing into one theme, so you read a topic not hundreds of headlines."),
        ("Spot what's rising",
         "It tracks each theme month to month and highlights the fastest-growing — the early warnings."),
        ("Hand off for review",
         "It presents ranked, plain summaries. A person reviews and decides — the tool never decides."),
    ]
    n = len(steps)
    gap = Inches(0.18)
    top = Inches(2.05)
    h = Inches(1.95)
    card_w = Emu(int((total_w - gap * (n - 1)) / n))
    strip_colors = [ACCENT, GREEN, ACCENT, YELLOW, CORE]
    for i, (title, desc) in enumerate(steps):
        x = Emu(int(margin + i * (card_w + gap)))
        _box(slide, x, top, card_w, h, fill=PANEL, line=LINE, line_w=1.0, rounded=True)
        disc = slide.shapes.add_shape(MSO_SHAPE.OVAL, x + Inches(0.15), top + Inches(0.15),
                                      Inches(0.42), Inches(0.42))
        disc.fill.solid()
        disc.fill.fore_color.rgb = CORE
        disc.line.fill.background()
        disc.shadow.inherit = False
        dtf = disc.text_frame
        dtf.margin_left = dtf.margin_right = dtf.margin_top = dtf.margin_bottom = Pt(0)
        dp = dtf.paragraphs[0]
        dp.alignment = PP_ALIGN.CENTER
        dr = dp.add_run()
        dr.text = str(i + 1)
        dr.font.size = Pt(16)
        dr.font.bold = True
        dr.font.color.rgb = WHITE
        dr.font.name = FONT
        _box(slide, x, top + h - Pt(5), card_w, Pt(5), fill=strip_colors[i % len(strip_colors)])
        _text(slide, x + Inches(0.15), top + Inches(0.66), card_w - Inches(0.3), Inches(0.5),
              [[(title, 12.5, CORE, True, False)]])
        _text(slide, x + Inches(0.15), top + Inches(1.02), card_w - Inches(0.3), Inches(0.85),
              [[(desc, 9.5, GRAY, False, False)]], line_spacing=1.08)

    band_top = Inches(4.35)
    band_h = Inches(1.2)
    half = Emu(int((total_w - Inches(0.3)) / 2))

    _box(slide, margin, band_top, half, band_h, fill=RGBColor(0xEE, 0xF3, 0xFA),
         line=LINE, line_w=1.0, rounded=True)
    _box(slide, margin, band_top, Inches(0.09), band_h, fill=ACCENT)
    _text(slide, margin + Inches(0.25), band_top + Inches(0.14), half - Inches(0.45), band_h - Inches(0.2),
          [[("A simple analogy", 12, CORE, True, False)],
           [("It's like a smoke detector for risk. It doesn't put out fires or decide what's "
             "dangerous — it notices smoke early and points you to the room to check.",
             10.5, BODY, False, False)]],
          line_spacing=1.12, space_after=3)

    lx = Emu(int(margin + half + Inches(0.3)))
    _box(slide, lx, band_top, half, band_h, fill=RGBColor(0xFF, 0xF9, 0xE6),
         line=LINE, line_w=1.0, rounded=True)
    _box(slide, lx, band_top, Inches(0.09), band_h, fill=YELLOW)
    _text(slide, lx + Inches(0.25), band_top + Inches(0.14), half - Inches(0.45), band_h - Inches(0.2),
          [[("What it can't do", 12, CORE, True, False)],
           [("It reads mostly headlines and summaries (not full articles), covers US/English "
             "news, and can miss nuance. A starting point for review, not the final word.",
             10.5, BODY, False, False)]],
          line_spacing=1.12, space_after=3)

    _text(slide, margin, Inches(5.75), total_w, Inches(0.4),
          [[("Same engine as the previous slide, described for a general audience. "
             "For definitions and common questions, see the FAQ page in the reports.",
             9.5, MUTED, False, True)]], line_spacing=1.0)


def build():
    prs = load_template()
    kept = len(prs.slides)
    build_technical_slide(prs)
    build_plain_language_slide(prs)
    OUT.parent.mkdir(exist_ok=True)
    prs.save(str(OUT))
    print(f"Wrote deck -> {OUT}  ({kept} template slide(s) + 2 new = {len(prs.slides)} total)")


if __name__ == "__main__":
    build()
